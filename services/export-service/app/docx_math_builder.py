"""
docx_math_builder.py
Native Office Math Markup Language (OMML) Document Builder for HSG Math Exam Studio.
Compatible with Microsoft Word Equation Editor without broken formatting, layout splits or font corruption.
"""

import os
import re
import tempfile
from typing import List, Dict, Any, Optional
import lxml.etree as ET
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from minio import Minio

def get_minio_client() -> Minio:
    endpoint = os.getenv("MINIO_ENDPOINT", "minio-storage:9000")
    access_key = os.getenv("MINIO_ROOT_USER", "minioadmin")
    secret_key = os.getenv("MINIO_ROOT_PASSWORD", "minioadmin")
    use_ssl = os.getenv("MINIO_USE_SSL", "false").lower() == "true"
    return Minio(endpoint=endpoint, access_key=access_key, secret_key=secret_key, secure=use_ssl)


class LaTeXMathConverter:
    """Preprocessor for Vietnamese Math LaTeX macros and MathML-to-OMML converter."""

    MML2OMML_XSL_FALLBACK = """<?xml version="1.0" encoding="UTF-8"?>
    <xsl:stylesheet version="1.0" 
        xmlns:xsl="http://www.w3.org/1999/XSL/Transform"
        xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
        xmlns:mml="http://www.w3.org/1998/Math/MathML">
        <xsl:output method="xml" encoding="UTF-8" indent="no"/>

        <xsl:template match="mml:math">
            <m:oMath>
                <xsl:apply-templates select="*"/>
            </m:oMath>
        </xsl:template>

        <xsl:template match="mml:mrow">
            <xsl:apply-templates select="*"/>
        </xsl:template>

        <xsl:template match="mml:mi | mml:mn">
            <m:r>
                <m:t><xsl:value-of select="."/></m:t>
            </m:r>
        </xsl:template>

        <xsl:template match="mml:mo">
            <m:r>
                <m:rPr><m:lit/></m:rPr>
                <m:t><xsl:value-of select="."/></m:t>
            </m:r>
        </xsl:template>

        <xsl:template match="mml:mfrac">
            <m:f>
                <m:num><xsl:apply-templates select="*[1]"/></m:num>
                <m:den><xsl:apply-templates select="*[2]"/></m:den>
            </m:f>
        </xsl:template>

        <xsl:template match="mml:msup">
            <m:sSup>
                <m:e><xsl:apply-templates select="*[1]"/></m:e>
                <m:sup><xsl:apply-templates select="*[2]"/></m:sup>
            </m:sSup>
        </xsl:template>

        <xsl:template match="mml:msub">
            <m:sSub>
                <m:e><xsl:apply-templates select="*[1]"/></m:e>
                <m:sub><xsl:apply-templates select="*[2]"/></m:sub>
            </m:sSub>
        </xsl:template>

        <xsl:template match="mml:msqrt">
            <m:rad>
                <m:radPr><m:degHide m:val="on"/></m:radPr>
                <m:deg/>
                <m:e><xsl:apply-templates select="*"/></m:e>
            </m:rad>
        </xsl:template>
    </xsl:stylesheet>
    """

    def __init__(self, xsl_path: Optional[str] = None):
        if xsl_path and os.path.exists(xsl_path):
            with open(xsl_path, "rb") as f:
                xslt_root = ET.XML(f.read())
        else:
            xslt_root = ET.XML(self.MML2OMML_XSL_FALLBACK.encode("utf-8"))
        self.transform = ET.XSLT(xslt_root)

    @staticmethod
    def preprocess_vietnamese_math_macros(latex_code: str) -> str:
        s = latex_code.strip()
        s = re.sub(r'\\heva\s*\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}', r'\\begin{cases} \1 \\end{cases}', s)
        s = re.sub(r'\\hoac\s*\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}', r'\\left[\\begin{array}{l} \1 \\end{array}\\right.', s)
        s = re.sub(r'\\widehat\s*\{([^}]+)\}', r'\\hat{\1}', s)
        s = re.sub(r'\\wideparen\s*\{([^}]+)\}', r'\\frown{\1}', s)
        s = s.replace(r'\dfrac', r'\frac')
        s = s.replace(r'\le ', r'\leq ').replace(r'\ge ', r'\geq ')
        s = s.replace(r'\not=', r'\neq ')
        s = s.replace(r'\parallel', r'\mathbin{\parallel}')
        s = s.replace(r'\perp', r'\bot')
        return s

    def latex_to_mathml(self, latex_str: str) -> str:
        clean_latex = self.preprocess_vietnamese_math_macros(latex_str)
        try:
            import latex2mathml.converter
            return latex2mathml.converter.convert(clean_latex)
        except Exception:
            return f'<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mi>{clean_latex}</mi></mrow></math>'

    def mathml_to_omml(self, mathml_str: str) -> ET._Element:
        mml_dom = ET.fromstring(mathml_str.encode("utf-8"))
        omml_dom = self.transform(mml_dom)
        return omml_dom.getroot()

    def convert_latex_to_omml_element(self, latex_str: str) -> ET._Element:
        mathml = self.latex_to_mathml(latex_str)
        return self.mathml_to_omml(mathml)


class DocxMathExamBuilder:
    """Document Builder producing 100% native OMML Word .docx matching MOET standards."""

    COLOR_PRIMARY = "1E3A8A"    # Indigo / Navy
    COLOR_BORDER = "94A3B8"     # Viền thanh mảnh Slate-400
    COLOR_ZEBRA = "F8FAFC"      # Slate-50 zebra background

    def __init__(self, math_converter: Optional[LaTeXMathConverter] = None):
        self.doc = docx.Document()
        self.math_converter = math_converter or LaTeXMathConverter()
        self._setup_page_geometry()
        self._setup_global_styles()

    def _setup_page_geometry(self):
        """Top/Bottom 1.8cm, Left 2.0cm, Right 1.8cm."""
        for section in self.doc.sections:
            section.top_margin = Inches(0.71)     # 1.8 cm
            section.bottom_margin = Inches(0.71)  # 1.8 cm
            section.left_margin = Inches(0.79)    # 2.0 cm
            section.right_margin = Inches(0.71)   # 1.8 cm
            section.page_width = Inches(8.27)     # A4
            section.page_height = Inches(11.69)   # A4

    def _setup_global_styles(self):
        normal_style = self.doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(12)
        normal_style.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_mixed_text_with_math(self, paragraph, text: str):
        """Splits $...$ or $$...$$ formulas and adds native <m:oMath> elements."""
        if not text:
            return
        tokens = re.split(r'(\$\$[^\$]+\$\$|\$[^\$]+\$)', text)
        for token in tokens:
            if not token:
                continue
            is_display = token.startswith('$$') and token.endswith('$$') and len(token) > 4
            is_inline = token.startswith('$') and token.endswith('$') and len(token) > 2

            if is_display or is_inline:
                latex_expr = token[2:-2] if is_display else token[1:-1]
                try:
                    omml_node = self.math_converter.convert_latex_to_omml_element(latex_expr)
                    xml_str = ET.tostring(omml_node, encoding="utf-8").decode("utf-8")
                    docx_omml = parse_xml(xml_str)
                    paragraph._p.append(docx_omml)
                except Exception:
                    run = paragraph.add_run(token)
                    run.italic = True
            else:
                paragraph.add_run(token)

    def add_exam_header(self, governing_body: str, school_name: str, exam_title: str, year_str: str, duration: int, exam_code: str = "ĐỀ CHÍNH THỨC"):
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._remove_table_borders(table)

        # Left Column
        c_left = table.rows[0].cells[0]
        c_left.width = Inches(3.6)
        p1 = c_left.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(governing_body.upper() + "\n")
        r1.font.size = Pt(10)
        r2 = p1.add_run(school_name.upper() + "\n")
        r2.font.bold = True
        r2.font.size = Pt(10.5)
        r3 = p1.add_run(f"{exam_code}\n")
        r3.font.bold = True
        r3.font.size = Pt(11)
        r4 = p1.add_run("(Đề thi gồm có 01 trang)")
        r4.font.italic = True
        r4.font.size = Pt(9.5)

        # Right Column
        c_right = table.rows[0].cells[1]
        c_right.width = Inches(3.5)
        p2 = c_right.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        r5 = p2.add_run(exam_title.upper() + "\n")
        r5.font.bold = True
        r5.font.size = Pt(11)
        r6 = p2.add_run(f"{year_str.upper()}\n")
        r6.font.bold = True
        r6.font.size = Pt(10.5)
        r7 = p2.add_run("Môn thi: TOÁN 9\n")
        r7.font.bold = True
        r7.font.size = Pt(10.5)
        r8 = p2.add_run(f"Thời gian làm bài: {duration} phút (không kể thời gian phát đề)")
        r8.font.italic = True
        r8.font.size = Pt(9.5)

        # Candidate Box
        p_box = self.doc.add_paragraph()
        p_box.paragraph_format.space_before = Pt(6)
        p_box.paragraph_format.space_after = Pt(8)
        table_sbd = self.doc.add_table(rows=1, cols=2)
        table_sbd.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._apply_table_borders(table_sbd)
        c0 = table_sbd.rows[0].cells[0]
        c0.width = Inches(4.8)
        c0.paragraphs[0].add_run("Họ và tên thí sinh: ................................................................").font.size = Pt(10)
        c1 = table_sbd.rows[0].cells[1]
        c1.width = Inches(2.3)
        c1.paragraphs[0].add_run("Số báo danh: ........................").font.size = Pt(10)

        # Divider
        p_div = self.doc.add_paragraph()
        p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_div.paragraph_format.space_before = Pt(4)
        p_div.paragraph_format.space_after = Pt(8)
        p_div.add_run("————————————————————————————").font.bold = True

    def add_question(self, q_title: str, q_content: str, points: float, sub_questions: Optional[List[Dict[str, Any]]] = None):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True

        r_title = p.add_run(f"{q_title} ({points:.1f} điểm): ")
        r_title.font.bold = True
        self.add_mixed_text_with_math(p, q_content)

        if sub_questions:
            for sq in sub_questions:
                p_sq = self.doc.add_paragraph()
                p_sq.paragraph_format.left_indent = Inches(0.25)
                p_sq.paragraph_format.space_before = Pt(2)
                p_sq.paragraph_format.space_after = Pt(2)
                p_sq.paragraph_format.line_spacing = 1.15

                sq_lbl = sq.get("label", "a)")
                sq_pts = sq.get("points")
                pts_str = f" ({sq_pts} điểm)" if sq_pts else ""
                r_sq = p_sq.add_run(f"{sq_lbl}{pts_str}: ")
                r_sq.font.bold = True
                self.add_mixed_text_with_math(p_sq, sq.get("content_latex", ""))

    def add_geometry_question_side_by_side(
        self,
        q_title: str,
        q_content: str,
        points: float,
        image_path: Optional[str] = None,
        sub_questions: Optional[List[Dict[str, Any]]] = None
    ):
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._remove_table_borders(table)

        row = table.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        cell_left = row.cells[0]
        cell_right = row.cells[1]
        cell_left.width = Inches(4.6)
        cell_right.width = Inches(2.5)
        cell_left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cell_right.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        p_left = cell_left.paragraphs[0]
        p_left.paragraph_format.line_spacing = 1.15
        p_left.paragraph_format.space_after = Pt(3)
        p_left.paragraph_format.keep_with_next = True
        r_title = p_left.add_run(f"{q_title} ({points:.1f} điểm): ")
        r_title.font.bold = True
        self.add_mixed_text_with_math(p_left, q_content)

        if sub_questions:
            for sq in sub_questions:
                p_sq = cell_left.add_paragraph()
                p_sq.paragraph_format.left_indent = Inches(0.15)
                p_sq.paragraph_format.space_after = Pt(2)
                sq_lbl = sq.get("label", "a)")
                sq_pts = sq.get("points")
                pts_str = f" ({sq_pts} điểm)" if sq_pts else ""
                r_sq = p_sq.add_run(f"{sq_lbl}{pts_str}: ")
                r_sq.font.bold = True
                self.add_mixed_text_with_math(p_sq, sq.get("content_latex", ""))

        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path and os.path.exists(image_path):
            r_img = p_right.add_run()
            r_img.add_picture(image_path, width=Inches(2.3))

    def add_marking_scheme_table(self, rubric_data: List[Dict[str, Any]]):
        p_head = self.doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(12)
        p_head.paragraph_format.space_after = Pt(6)
        p_head.paragraph_format.keep_with_next = True
        r_head = p_head.add_run("HƯỚNG DẪN CHẤM VÀ THANG ĐIỂM CHI TIẾT (CHUẨN BỘ GD&ĐT)")
        r_head.font.bold = True
        r_head.font.size = Pt(12.5)

        table = self.doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_widths_inch = [Inches(1.0), Inches(5.3), Inches(0.8)]
        col_widths_dxa = [1417, 7511, 1134]

        hdr_row = table.rows[0]
        hdr_tr = hdr_row._tr.get_or_add_trPr()
        hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        hdr_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        headers_text = ["Câu / Ý", "Nội dung lập luận & Các bước giải chi tiết", "Điểm"]
        for idx, text in enumerate(headers_text):
            cell = hdr_row.cells[idx]
            cell.width = col_widths_inch[idx]
            self._set_cell_width_dxa(cell, col_widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{self.COLOR_PRIMARY}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10.5)

        for row_idx, item in enumerate(rubric_data):
            row = table.add_row()
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            c0 = row.cells[0]
            c0.width = col_widths_inch[0]
            self._set_cell_width_dxa(c0, col_widths_dxa[0])
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(item.get("question_label", ""))
            r0.font.bold = True
            r0.font.size = Pt(10.5)

            c1 = row.cells[1]
            c1.width = col_widths_inch[1]
            self._set_cell_width_dxa(c1, col_widths_dxa[1])
            c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p1 = c1.paragraphs[0]
            p1.paragraph_format.line_spacing = 1.15
            p1.paragraph_format.space_before = Pt(3)
            p1.paragraph_format.space_after = Pt(3)
            self.add_mixed_text_with_math(p1, item.get("solution_step", ""))

            c2 = row.cells[2]
            c2.width = col_widths_inch[2]
            self._set_cell_width_dxa(c2, col_widths_dxa[2])
            c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            score_val = item.get("score", 0.25)
            r2 = p2.add_run(f"{score_val:.2f}" if isinstance(score_val, float) else str(score_val))
            r2.font.bold = True
            r2.font.size = Pt(10.5)

            if row_idx % 2 == 1:
                for c in [c0, c1, c2]:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{self.COLOR_ZEBRA}"/>')
                    c._tc.get_or_add_tcPr().append(shd)

        self._apply_table_borders(table)

    def add_fixed_layout_rubric_table(self, essay_questions: List[Dict[str, Any]]):
        """
        Dựng bảng 3 cột Hướng dẫn chấm theo kỹ thuật Sub-Row + vMerge (gộp ô dọc) cho cột Câu.
        Điểm số luôn căn ngang tuyệt đối với dòng biến đổi tương ứng.
        """
        p_head = self.doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(12)
        p_head.paragraph_format.space_after = Pt(6)
        p_head.paragraph_format.keep_with_next = True
        r_head = p_head.add_run("HƯỚNG DẪN CHẤM VÀ THANG ĐIỂM CHI TIẾT (CHUẨN BỘ GD&ĐT)")
        r_head.font.bold = True
        r_head.font.size = Pt(12.5)

        table = self.doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_widths_inch = [Inches(1.0), Inches(5.3), Inches(0.8)]
        col_widths_dxa = [1417, 7511, 1134]

        # Header row
        hdr_row = table.rows[0]
        hdr_tr = hdr_row._tr.get_or_add_trPr()
        hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        hdr_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        headers_text = ["Câu", "Nội dung lập luận & Các bước giải chi tiết", "Điểm"]
        for idx, text in enumerate(headers_text):
            cell = hdr_row.cells[idx]
            cell.width = col_widths_inch[idx]
            self._set_cell_width_dxa(cell, col_widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{self.COLOR_PRIMARY}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10.5)

        for q in essay_questions:
            sub_label = q.get("sub_label", "")
            prompt = q.get("prompt_latex", "")
            steps = q.get("grading_steps", [])

            # Hàng 0: Đề bài
            row_prompt = table.add_row()
            trPr0 = row_prompt._tr.get_or_add_trPr()
            trPr0.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            c0 = row_prompt.cells[0]
            self._set_cell_width_dxa(c0, col_widths_dxa[0])
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcPr0 = c0._tc.get_or_add_tcPr()
            tcPr0.append(parse_xml(f'<w:vMerge {nsdecls("w")} w:val="restart"/>'))
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for line in sub_label.split("\n"):
                r = p0.add_run(line + "\n")
                r.font.bold = True
                r.font.size = Pt(10.5)

            c1 = row_prompt.cells[1]
            self._set_cell_width_dxa(c1, col_widths_dxa[1])
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_before = Pt(3)
            p1.paragraph_format.space_after = Pt(3)
            self.add_mixed_text_with_math(p1, prompt)
            for r in p1.runs:
                r.font.bold = True

            c2 = row_prompt.cells[2]
            self._set_cell_width_dxa(c2, col_widths_dxa[2])

            # Các hàng bước giải
            for s_idx, step in enumerate(steps):
                is_last = (s_idx == len(steps) - 1)
                row_step = table.add_row()
                trPr_s = row_step._tr.get_or_add_trPr()
                trPr_s.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

                # Cột 0: vMerge CONTINUE
                c0s = row_step.cells[0]
                self._set_cell_width_dxa(c0s, col_widths_dxa[0])
                tcPr0s = c0s._tc.get_or_add_tcPr()
                tcPr0s.append(parse_xml(f'<w:vMerge {nsdecls("w")}/>'))

                # Cột 1: Bước giải
                c1s = row_step.cells[1]
                self._set_cell_width_dxa(c1s, col_widths_dxa[1])
                p1s = c1s.paragraphs[0]
                p1s.paragraph_format.line_spacing = 1.15
                p1s.paragraph_format.space_before = Pt(2)
                p1s.paragraph_format.space_after = Pt(2)
                content = step.get("latex_content", "")
                self.add_mixed_text_with_math(p1s, content)

                # Ẩn border ngang giữa các bước, chỉ đóng ở bước cuối
                border_bottom = 'w:val="single" w:sz="4" w:space="0" w:color="94A3B8"' if is_last else 'w:val="nil"'
                border_top = 'w:val="nil"'

                for cell in [c0s, c1s]:
                    borders = parse_xml(f'''
                        <w:tcBorders {nsdecls("w")}>
                            <w:top {border_top}/>
                            <w:bottom {border_bottom}/>
                        </w:tcBorders>
                    ''')
                    cell._tc.get_or_add_tcPr().append(borders)

                # Cột 2: Điểm
                c2s = row_step.cells[2]
                self._set_cell_width_dxa(c2s, col_widths_dxa[2])
                c2s.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p2s = c2s.paragraphs[0]
                p2s.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pts = step.get("points")
                if pts is not None:
                    r2s = p2s.add_run(f"{pts:.2f}")
                    r2s.font.bold = True
                    r2s.font.size = Pt(10.5)

                borders2 = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:top {border_top}/>
                        <w:bottom {border_bottom}/>
                    </w:tcBorders>
                ''')
                c2s._tc.get_or_add_tcPr().append(borders2)

        self._apply_table_borders(table)

    @staticmethod
    def _set_cell_width_dxa(cell, dxa_width: int):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{dxa_width}" w:type="dxa"/>')
        tcPr.append(tcW)

    @staticmethod
    def _remove_table_borders(table):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(f"""
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        """)
        tblPr.append(tblBorders)

    def _apply_table_borders(self, table):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(f"""
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{self.COLOR_BORDER}"/>
            </w:tblBorders>
        """)
        tblPr.append(tblBorders)

    def save(self, filepath: str):
        self.doc.save(filepath)
        return filepath


def compile_native_omml_docx(exam_data: Dict[str, Any], mode: str = "EXAM") -> str:
    """Compiles complete exam into native OMML Word .docx and uploads to MinIO."""
    builder = DocxMathExamBuilder()
    exam_id = str(exam_data.get("id", "exam"))
    title = exam_data.get("title", "KỲ THI CHỌN HỌC SINH GIỎI TOÁN 9")
    header = exam_data.get("header_config") or {}
    slots = exam_data.get("slots_state") or []

    minio_client = get_minio_client()
    bucket_name = os.getenv("MINIO_BUCKET_NAME", "hsg-math-assets")

    with tempfile.TemporaryDirectory() as td:
        # Download figures
        local_figures = {}
        for slot in slots:
            q = slot.get("question") or {}
            figs = q.get("figures_meta") or []
            for fig in figs:
                f_key = fig.get("figure_key")
                if f_key:
                    local_p = os.path.join(td, os.path.basename(f_key))
                    try:
                        minio_client.fget_object(bucket_name, f_key, local_p)
                        local_figures[f_key] = local_p
                    except Exception as e:
                        print(f"[DocxBuilder Warning] Could not fetch figure {f_key}: {e}")

        # Add Header
        gov_body = header.get("governing_body", "SỞ GIÁO DỤC VÀ ĐÀO TẠO")
        school = header.get("department_or_school", "TRƯỜNG THCS CHUYÊN LÊ QUÝ ĐÔN")
        exam_name = header.get("exam_name", title)
        acad_year = header.get("academic_year", "Năm học 2026 - 2027")
        duration = header.get("duration_minutes", 150)
        code = header.get("exam_code", "ĐỀ CHÍNH THỨC")

        builder.add_exam_header(gov_body, school, exam_name, acad_year, duration, code)

        if mode.upper() == "EXAM":
            # Questions
            for idx, slot in enumerate(slots):
                q_num = f"Câu {slot.get('slot_index', idx + 1)}"
                pts = float(slot.get("points") or 1.0)
                q = slot.get("question") or {}
                content = q.get("content_latex", "")
                sub_qs = q.get("sub_questions") or []
                figs = q.get("figures_meta") or []
                topic = q.get("main_topic", "")

                fig_path = None
                if figs and figs[0].get("figure_key") in local_figures:
                    fig_path = local_figures[figs[0].get("figure_key")]

                if topic == "PLANE_GEOMETRY" and fig_path:
                    builder.add_geometry_question_side_by_side(q_num, content, pts, fig_path, sub_qs)
                else:
                    builder.add_question(q_num, content, pts, sub_qs)

            # Footer
            p_end = builder.doc.add_paragraph()
            p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_end.paragraph_format.space_before = Pt(16)
            r_end = p_end.add_run("--- HẾT ---\nCán bộ coi thi không giải thích gì thêm.")
            r_end.font.italic = True
            r_end.font.size = Pt(10)
        else:
            # Rubric Table
            template_doc = exam_data.get("template_document")
            essay_questions = template_doc.get("part2_essay", {}).get("questions", []) if template_doc else []

            if essay_questions:
                builder.add_fixed_layout_rubric_table(essay_questions)
            else:
                rubric_rows = []
                for idx, slot in enumerate(slots):
                    q_lbl = f"Câu {slot.get('slot_index', idx + 1)}"
                    pts = float(slot.get("points") or 1.0)
                    q = slot.get("question") or {}
                    rubric = slot.get("rubric_data") or {}
                    steps = rubric.get("steps") or []

                    if steps:
                        for s_idx, st in enumerate(steps):
                            tag = st.get("milestone_tag", "")
                            desc = f"{tag}: {st.get('latex_content', '')}" if tag else st.get('latex_content', '')
                            rubric_rows.append({
                                "question_label": q_lbl if s_idx == 0 else "",
                                "solution_step": desc,
                                "score": st.get("points", 0.25)
                            })
                    else:
                        sol = q.get("solution_latex", "*(Đang cập nhật lời giải)*")
                        rubric_rows.append({
                            "question_label": q_lbl,
                            "solution_step": sol,
                            "score": pts
                        })

                builder.add_marking_scheme_table(rubric_rows)

        out_docx = os.path.join(td, "exam.docx")
        builder.save(out_docx)

        s3_key = f"exports/{exam_id}/native_omml_{mode.lower()}.docx"
        minio_client.fput_object(
            bucket_name,
            s3_key,
            out_docx,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        return s3_key
